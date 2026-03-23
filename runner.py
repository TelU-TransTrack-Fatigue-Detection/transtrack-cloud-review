import asyncio
import logging
import random
import time
from datetime import datetime, timedelta
from pathlib import Path

import httpx
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from config import settings

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger(__name__)

OUTPUT_DIR = Path("output")
EXCEL_PATH = OUTPUT_DIR / "results.xlsx"
TMP_DIR    = Path("records/runner_tmp")
POLL_EVERY = 60
LOOKBACK_HOURS = 2

EXCEL_HEADERS = [
    "event_id", "imei", "time", "alarm",
    "confidence_level", "review_result", "process_duration_ms",
    "video_url", "processed_at",
]


def login() -> tuple[str, str]:
    resp = requests.post(
        f"{settings.TRANSTRACK_BASE_URL}/auth",
        json={"username": settings.TRANSTRACK_USERNAME, "password": settings.TRANSTRACK_PASSWORD},
        timeout=30,
    )
    resp.raise_for_status()
    data = resp.json().get("data", {})
    token, pid = data.get("token"), data.get("pid")
    if not token or not pid:
        raise RuntimeError(f"Login failed: {resp.json()}")
    logger.info("Logged in")
    return token, pid


def make_session(token: str, pid: str) -> requests.Session:
    session = requests.Session()
    session.headers.update({
        "X-Token": token,
        "X-User": pid,
        "Authorization": f"Bearer {token}",
    })
    return session


def fetch_events(session: requests.Session) -> list[dict]:
    now = datetime.now()
    date_start = (now - timedelta(hours=LOOKBACK_HOURS)).strftime("%Y-%m-%d %H:%M:%S").replace(" ", "%20")
    date_end   = now.strftime("%Y-%m-%d %H:%M:%S").replace(" ", "%20")

    url = (
        f"{settings.TRANSTRACK_BASE_URL}/events"
        f"?order=desc&sort=device_time"
        f"&filter_columns=alarm_type&filter_value=121,122"
        f"&page=1&page_size=50"
        f"&range_date_columns=device_time"
        f"&range_date_start={date_start}&range_date_end={date_end}"
    )

    resp = session.get(url, timeout=180)
    resp.raise_for_status()
    return resp.json().get("data", {}).get("list", [])


async def download_video(url: str, path: Path) -> bool:
    try:
        async with httpx.AsyncClient(timeout=settings.VIDEO_DOWNLOAD_TIMEOUT) as client:
            async with client.stream("GET", url) as resp:
                resp.raise_for_status()
                with open(path, "wb") as f:
                    async for chunk in resp.aiter_bytes(chunk_size=1024 * 64):
                        f.write(chunk)
        if path.stat().st_size < 3072:
            path.unlink(missing_ok=True)
            return False
        return True
    except Exception as e:
        logger.warning("Download failed: %s", e)
        path.unlink(missing_ok=True)
        return False


async def run_inference(video_path: str, alarm: str) -> dict:
    await asyncio.sleep(random.uniform(0.2, 0.8))
    confidence_level = random.randint(40, 95)
    return {
        "confidence_level": confidence_level,
        "review_result": confidence_level >= 60,
    }


def init_excel():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if EXCEL_PATH.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(EXCEL_PATH)
    logger.info("Created Excel at %s", EXCEL_PATH)


def append_excel(row: dict):
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    ws.append([row.get(h, "") for h in EXCEL_HEADERS])
    wb.save(EXCEL_PATH)


async def process_event(event: dict, video_url: str) -> dict | None:
    alarm_id = event.get("identity", event.get("id", ""))
    alarm    = "eyes_closed" if event.get("alarm_type") == "121" else "yawning"
    imei     = event.get("device", {}).get("imei", "")
    time_str = event.get("time", "")

    video_path = TMP_DIR / f"{alarm_id}.mp4"
    start = time.monotonic()

    ok = await download_video(video_url, video_path)
    if not ok:
        logger.warning("Skipping id=%s — bad video file", alarm_id)
        return None

    inference = await run_inference(str(video_path), alarm)
    process_duration = int((time.monotonic() - start) * 1000)

    video_path.unlink(missing_ok=True)

    result = {
        "event_id":            alarm_id,
        "imei":                imei,
        "time":                time_str,
        "alarm":               alarm,
        "confidence_level":    inference["confidence_level"],
        "review_result":       inference["review_result"],
        "process_duration_ms": process_duration,
        "video_url":           video_url,
        "processed_at":        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    append_excel(result)
    logger.info(
        "Processed id=%s alarm=%s confidence=%d result=%s duration=%dms",
        alarm_id, alarm, result["confidence_level"], result["review_result"], process_duration,
    )
    return result


async def main():
    if not settings.TRANSTRACK_BASE_URL or not settings.TRANSTRACK_USERNAME:
        logger.error("Fill in TRANSTRACK_BASE_URL, TRANSTRACK_USERNAME, TRANSTRACK_PASSWORD in .env")
        return

    TMP_DIR.mkdir(parents=True, exist_ok=True)
    init_excel()

    token, pid = login()
    session = make_session(token, pid)
    seen_ids: set[str] = set()

    logger.info("Runner started — polling every %ds", POLL_EVERY)

    while True:
        try:
            events = fetch_events(session)
            new_events = [
                e for e in events
                if (e.get("identity") or e.get("id")) not in seen_ids
            ]
            logger.info("Fetched %d events, %d new", len(events), len(new_events))

            for event in new_events:
                event_id = event.get("identity") or event.get("id")
                seen_ids.add(event_id)

                alarm_files = event.get("alarm_file", [])
                if not alarm_files:
                    continue

                video_url = alarm_files[0].get("downUrl", "").replace("&amp;", "&")
                if not video_url:
                    continue

                await process_event(event, video_url)

        except Exception as e:
            logger.error("Poll error: %s", e)

        await asyncio.sleep(POLL_EVERY)


if __name__ == "__main__":
    asyncio.run(main())
